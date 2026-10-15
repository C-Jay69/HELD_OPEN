# -*- coding: utf-8 -*-
"""Apply the D3 revision edits, emit the revised manuscript and the change log."""
import json, re, sys, os, difflib
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from edits import EDITS, DIVIDER_FROM, DIVIDER_TO

SRC = os.path.join(HERE, "d3_paras.json")
OUT = os.path.abspath(os.path.join(HERE, ".."))
os.makedirs(OUT, exist_ok=True)

paras = json.loads(open(SRC).read())
orig = list(paras)

def words(txt):
    return len(txt.split())

w_before = words(" ".join(orig))

# --- section map -------------------------------------------------------
HEAD = re.compile(r"^(Chapter (One|Two|Three|Four|Five|Six|Seven|Eight|Nine|Ten|Eleven|Twelve|Thirteen|Fourteen|Fifteen|Sixteen|Seventeen|Eighteen|Nineteen|Twenty|Twenty-One|Twenty-Two|Twenty-Three|Twenty-Four|Twenty-Five|Twenty-Six|Twenty-Seven|Twenty-Eight|Twenty-Nine|Thirty|Thirty-One|Thirty-Two|Thirty-Three|Thirty-Four) *[-\u2013\u2014:] .+|CUSTOS Interlude I: Dakar|CUSTOS Interlude II: Notices of the Office of Records and Accommodation|CUSTOS Interlude III: Kestrel Field|CUSTOS Interlude IV: Held Open|CUSTOS Interlude IV|CUSTOS Section: Kestrel Field|Notices of the Office of Records and Accommodation|Coda *[-\u2013\u2014] .+)$")
sec_of = {}
sections = []
for i, t in enumerate(paras):
    s = t.strip()
    if HEAD.match(s):
        sections.append({"title": s, "start": i, "paras": []})
    if sections:
        sec_of[i] = sections[-1]["title"]
        if i != sections[-1]["start"]:
            sections[-1]["paras"].append(i)

def loc(idx):
    return sec_of.get(idx, "front matter")

# --- apply edits -------------------------------------------------------
def resolve(hint, needle):
    """Find the paragraph the edit belongs to. The edit list carries the index from
    the extraction it was built on; paragraphing can differ by a few indices between
    equally valid reconstructions of the same PDF, so fall back to a document-wide
    search, preferring the candidate nearest the hint."""
    if hint < len(paras) and needle in paras[hint]:
        return hint, "index"
    hits = [i for i, t in enumerate(paras) if needle in t]
    if len(hits) == 1:
        return hits[0], "unique"
    if hits:
        return min(hits, key=lambda i: abs(i - hint)), "nearest"
    return None, "none"

applied = []
errors = []
for eid, cat, idx_hint, old, new, note in EDITS:
    idx, how = resolve(idx_hint, old)
    if idx is None:
        errors.append((eid, idx_hint, 0, old[:60]))
        continue
    t = paras[idx]
    n = t.count(old)
    if n != 1:
        errors.append((eid, idx, n, old[:60]))
        continue
    paras[idx] = t.replace(old, new)
    applied.append((eid, cat, idx, old, new, note))
    if how != "index":
        print(f"  {eid}: anchored by {how} search -> paragraph {idx}")

if errors:
    print("FAILED EDITS:")
    for e in errors:
        print("  ", e)
    sys.exit(1)

# --- divider normalisation -------------------------------------------
div_count = 0
div_idx = []
for i, t in enumerate(paras):
    if t.strip() == DIVIDER_FROM:
        paras[i] = DIVIDER_TO
        div_count += 1
        div_idx.append(i)

w_after = words(" ".join(paras))

def body_count(paragraphs):
    start = next(i for i, t in enumerate(paragraphs)
                 if i >= 42 and t.strip().startswith("Chapter One"))
    n = 0
    for t in paragraphs[start:]:
        s = t.strip()
        n += 1 if s in ("---", "\u2022 \u2022 \u2022") else words(t)
    return n
print(f"edits applied: {len(applied)}  dividers normalised: {div_count}")
print(f"words: {w_before} -> {w_after}  ({w_after - w_before:+d})")

# --- deliverables ------------------------------------------------------
body_start = next(i for i, t in enumerate(paras) if i >= 42 and HEAD.match(t.strip()))
front = [t for t in paras[:body_start] if t.strip()]
toc = [t.strip() for t in front if t.strip() not in ("THE LAST SMUGGLERS", "SIMON J CLEARY", "Contents")]

# 1. markdown
md = ["# THE LAST SMUGGLERS", "", "*Simon J Cleary*", "", "## Contents", ""]
md += [f"- {t}" for t in toc]
md += ["", "---", ""]
for i in range(body_start, len(paras)):
    t = paras[i].strip()
    if not t:
        continue
    if HEAD.match(t):
        md += ["", f"## {t}", ""]
    elif t == DIVIDER_TO:
        md += ["", "*\u2022 \u2022 \u2022*", ""]
    elif re.match(r"^(Day \d+|Days \d+[\u2013-]\d+)", t):
        md += [f"*{t}*", ""]
    else:
        md += [t, ""]
open(f"{OUT}/D3_REVISED_manuscript.md", "w").write("\n".join(md))

# 2. docx (clean) and 3. docx (marked)
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX

EDIT_BY_IDX = {a[2]: a for a in applied}

def write_docx(path, marked):
    global toc
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(11)
    p = d.add_paragraph()
    r = p.add_run("THE LAST SMUGGLERS")
    r.bold = True
    r.font.size = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph()
    r = p.add_run("Simon J Cleary")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # front matter: epigraph page (the book had none; see SUBMISSION-READINESS.md)
    if marked:
        p = d.add_paragraph("EPIGRAPH — for decision; see SUBMISSION-READINESS.md §2")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep = d.add_paragraph()
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep.add_run("Nothing in this notice requires a response.").italic = True
    ep.add_run().add_break()
    ep.add_run("There is nothing you need to do.").italic = True
    d.add_page_break()
    d.add_heading("Contents", level=1)
    for t in toc:
        d.add_paragraph(t)
    d.add_page_break()
    for i in range(body_start, len(paras)):
        t = paras[i].strip()
        if not t and not (marked and i in EDIT_BY_IDX):
            continue
        if HEAD.match(t):
            d.add_heading(t, level=1)
            continue
        if t == DIVIDER_TO:
            p = d.add_paragraph("\u2022 \u2022 \u2022")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        p = d.add_paragraph()
        if re.match(r"^(Day \d+|Days \d+[\u2013-]\d+)", t):
            p.runs and None
            r = p.add_run(t)
            r.italic = True
            continue
        if marked and i in EDIT_BY_IDX:
            eid, cat, idx, old, new, note = EDIT_BY_IDX[i]
            sm = difflib.SequenceMatcher(None, old, new, autojunk=False)
            cut = [o for o in sm.get_opcodes() if o[0] != "equal"]
            if cut:
                a1, a2, b1, b2 = cut[0][1], cut[-1][2], cut[0][3], cut[-1][4]
            else:
                a1 = a2 = b1 = b2 = 0
            before = old[:a1]
            removed = old[a1:a2]
            inserted = new[b1:b2]
            after = new[b2:]
            op = orig[i]
            ip = op.index(old)
            pre = op[:ip]
            post = op[ip + len(old):]
            if pre:
                p.add_run(pre)
            if before:
                p.add_run(before)
            if removed:
                r = p.add_run(removed)
                r.font.strike = True
                r.font.color.rgb = docx.shared.RGBColor(0xC0, 0x00, 0x00)
            if inserted:
                r = p.add_run(inserted)
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if after:
                p.add_run(after)
            if post:
                p.add_run(post)
        else:
            p.add_run(t)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("END")
    d.save(path)

write_docx(f"{OUT}/D3_REVISED_manuscript.docx", False)
write_docx(f"{OUT}/D3_REVISED_marked.docx", True)

# 4. change log data for the writer script
_d = docx.Document(f"{OUT}/D3_REVISED_manuscript.docx")
docx_words = sum(len(p.text.split()) for p in _d.paragraphs)
edited_paras = sorted({a[2] for a in applied})
divider_paras = sorted(set(div_idx))
payload = {
    "words_before": w_before, "words_after": w_after,
    "body_before": body_count(orig), "body_after": body_count(paras),
    "docx_words": docx_words,
    "prose_delta": sum(words(a[4]) - words(a[3]) for a in applied),
    "edits_paras": len(edited_paras), "divider_paras": len(divider_paras),
    "untouched": len(paras) - len(set(edited_paras) | set(divider_paras)),
    "applied": [{"id": a[0], "cat": a[1], "idx": a[2], "loc": loc(a[2]),
                 "old": a[3], "new": a[4], "note": a[5]} for a in applied],
    "dividers": div_count,
    "sections": len([s for s in sections]),
}
json.dump(payload, open(os.path.join(HERE, "applied.json"), "w"), indent=1)
print("written:", OUT)
